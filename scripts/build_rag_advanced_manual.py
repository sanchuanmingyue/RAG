from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "storage" / "exports"
OUT_PATH = OUT_DIR / f"paper_reader_rag_advanced_manual_{datetime.now():%Y%m%d}.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
BORDER = "D0D5DD"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_font(run, name="Calibri", east_asia="Microsoft YaHei") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("PaperReader-RAG Advanced Guide")
    set_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("科研论文阅读 RAG 系统进阶构建方案与面试知识点手册")
    set_font(run)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("覆盖章节切分、父子分块、混合检索、RRF、Reranker、查询改写/意图分类、引用约束与拒答")
    set_font(r)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "交付说明",
        "本文档对应当前项目实现，而不是抽象方案。普通 RAG 仍使用 Chroma 与 Streamlit，不引入 LangChain、Milvus 或 PostgreSQL 的重构；新增能力以轻量模块形式接入现有链路。",
    )


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    set_font(r2)
    r2.font.size = Pt(10.5)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_font(r)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_font(r)
            r.font.size = Pt(9.5)
    doc.add_paragraph()


def add_module_section(doc: Document, title: str, implemented: list[str], knowledge: list[str], questions: list[tuple[str, str]]) -> None:
    doc.add_heading(title, level=2)
    doc.add_heading("当前实现", level=3)
    add_bullets(doc, implemented)
    doc.add_heading("核心知识点", level=3)
    add_bullets(doc, knowledge)
    doc.add_heading("面试官可能追问", level=3)
    for q, a in questions:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"问题：{q}")
        set_font(r)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(f"回答：{a}")
        set_font(r2)


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("一、项目总定位", level=1)
    doc.add_paragraph(
        "本项目面向科研论文阅读、文献综述和论文汇报场景，在已有普通 RAG 链路上新增检索质量优化与可信回答机制。"
        "系统目标不是简单回答 PDF 内容，而是让回答具备结构化检索、章节定位、证据引用和无依据拒答能力。"
    )
    add_table(
        doc,
        ["能力", "当前项目实现", "代码入口"],
        [
            ["章节切分", "识别论文标题并写入 section metadata", "backend/text_splitter.py"],
            ["父子分块", "子 chunk 入库检索，命中后按 parent_id 展开上下文", "backend/text_splitter.py / backend/vector_store.py"],
            ["混合检索", "Chroma 向量检索 + 本地关键词 BM25-like 检索", "backend/vector_store.py"],
            ["RRF 融合", "向量召回与关键词召回按 Reciprocal Rank Fusion 合并", "backend/vector_store.py"],
            ["Reranker", "轻量词项重叠、短语命中、RRF、距离综合排序", "backend/vector_store.py"],
            ["查询改写/意图分类", "规则识别实验、方法、相关工作、贡献类问题并扩展检索词", "backend/rag_chain.py"],
            ["引用约束与拒答", "Prompt 强制 [S1] 引用，生成后无引用则保守拒答", "backend/prompts.py / backend/rag_chain.py"],
        ],
        [1.15, 3.65, 1.7],
    )

    doc.add_heading("二、系统整体架构", level=1)
    add_numbered(
        doc,
        [
            "用户上传 PDF，系统按页提取文本并保留 paper_id、file_name、page。",
            "文本切分器先识别章节，再在章节 block 内执行滑动窗口子 chunk 切分。",
            "每个子 chunk 写入 Chroma，metadata 中保存章节、父块和来源字段。",
            "用户提问后，RAG 层执行意图分类与查询改写。",
            "向量检索与关键词检索并行召回候选结果。",
            "RRF 将多路召回结果融合，轻量 reranker 重新排序。",
            "命中子 chunk 后按 parent_id 展开父上下文，构造带来源编号的 Prompt。",
            "LLM 必须基于来源回答；无来源或无引用时触发拒答策略。",
        ],
    )

    add_callout(
        doc,
        "重建索引提醒",
        "新增 parent_id、child_index、section_type、section_path 等 metadata 后，旧向量库中的旧 chunk 不具备这些字段。上线后需要清空 Chroma collection 并重新上传论文。"
    )

    doc.add_heading("三、模块设计与知识点", level=1)
    add_module_section(
        doc,
        "3.1 章节切分",
        [
            "支持 Abstract、Introduction、Related Work、Method/Approach、Experiments/Evaluation、Results/Ablation、Conclusion、References。",
            "章节标题可带数字编号，例如 1 Introduction、2.1 Ablation Study。",
            "章节跨页时继承 current_section 状态，避免后一页 chunk 丢失章节归属。",
        ],
        [
            "章节 metadata 是论文 RAG 的重要结构信号，可以帮助按问题类型缩小检索范围。",
            "章节识别不替代滑动窗口，而是先确定结构边界，再在结构内部切分。",
        ],
        [
            ("为什么按章节切分比只按固定长度更适合论文？", "论文问题天然和章节相关，方法问题多在 Method，实验问题多在 Experiments/Results。章节 metadata 能提高召回精度，也能让回答来源更清楚。"),
            ("章节识别失败怎么办？", "系统仍会保留 page、chunk_id 等基础 metadata，并退化为普通滑动窗口检索；章节过滤无结果时也会回退全文检索。"),
        ],
    )

    add_module_section(
        doc,
        "3.2 父子分块",
        [
            "父块是章节 block 或页内章节 block，子块是用于 embedding 与检索的滑动窗口 chunk。",
            "子 chunk metadata 中保存 parent_id、parent_index、child_index。",
            "检索命中子块后，可按 parent_id 取回同父块内的兄弟 chunk，合并为更完整上下文。",
        ],
        [
            "子块小，适合精确召回；父块大，适合回答时保留上下文完整性。",
            "父子分块解决了“检索要精确、生成要完整”的冲突。",
        ],
        [
            ("为什么不直接索引父块？", "父块过大时 embedding 语义会变稀释，召回精度下降。索引子块可以更准确命中问题相关位置。"),
            ("父块展开会不会引入噪声？", "会，所以父块应控制在章节或页内章节范围内，而不是整篇论文。当前实现默认可通过 ENABLE_PARENT_CONTEXT 控制。"),
        ],
    )

    add_module_section(
        doc,
        "3.3 混合检索与 RRF",
        [
            "向量检索负责语义相似，关键词检索负责术语、缩写、指标和数据集名称命中。",
            "关键词检索使用本地 BM25-like 打分，不引入额外服务。",
            "RRF 将 vector rank 与 keyword rank 按 1 / (k + rank) 融合。",
        ],
        [
            "向量检索对同义表达友好，但可能漏掉精确术语。",
            "关键词检索对 exact match 友好，但不理解语义。",
            "RRF 不依赖不同检索器分数尺度，适合融合异构召回。",
        ],
        [
            ("为什么不用简单加权分数？", "向量距离和关键词分数尺度不同，直接相加不稳定。RRF 基于排名融合，更鲁棒。"),
            ("混合检索适合哪些论文问题？", "适合包含专有名词、baseline、dataset、算法缩写、指标名称的问题。"),
        ],
    )

    add_module_section(
        doc,
        "3.4 Reranker 精排",
        [
            "当前实现为轻量本地 reranker，不额外调用外部模型。",
            "综合考虑 query-token overlap、短语命中、RRF 分数、关键词分数和向量距离。",
            "可通过 ENABLE_RERANKER 关闭。",
        ],
        [
            "召回阶段关注覆盖率，精排阶段关注最终 top-k 的准确性。",
            "真实生产系统可替换为 cross-encoder 或 API reranker，但本项目保留轻量实现便于学习和演示。",
        ],
        [
            ("Reranker 和 RRF 有什么区别？", "RRF 是多路召回结果融合，解决候选集合合并；Reranker 是对候选集合重新排序，解决最终上下文质量。"),
            ("轻量 reranker 的局限是什么？", "它主要基于词项和简单信号，无法像 cross-encoder 那样深入判断问答相关性。"),
        ],
    )

    add_module_section(
        doc,
        "3.5 查询改写与意图分类",
        [
            "analyze_question() 将问题识别为 experiment、method、related_work、contribution 或 general。",
            "不同意图映射到不同 section_types。",
            "改写后的检索 query 会追加实验、方法、相关工作等英文提示词，提高召回稳定性。",
        ],
        [
            "意图分类可以减少无关章节干扰。",
            "查询改写可以把口语问题扩展成更适合检索的表达。",
            "章节过滤无结果时必须回退全文检索，避免过度过滤。"
        ],
        [
            ("为什么查询改写不用每次调用 LLM？", "规则改写成本低、稳定、可解释，适合课程项目和面试展示。后续可以加 LLM rewrite 作为增强模式。"),
            ("意图分类错了怎么办？", "系统会在章节过滤没有命中时回退全文检索，保证召回不被完全截断。"),
        ],
    )

    add_module_section(
        doc,
        "3.6 引用约束与拒答",
        [
            "Prompt 要求每个关键结论必须引用 [S1]、[S2]。",
            "如果检索不到来源，直接返回“论文中没有找到明确依据。”。",
            "STRICT_CITATION=true 时，如果模型回答没有引用标记，会触发保守拒答。",
        ],
        [
            "RAG 不能完全消除幻觉，必须通过来源约束和后处理降低风险。",
            "科研论文阅读场景比闲聊更强调可核查性，宁可拒答也不要编造。",
        ],
        [
            ("为什么有 sources 还要检查回答里有没有引用？", "因为模型可能忽略指令，给出无来源结论。引用检查是最后一道廉价防线。"),
            ("拒答会不会降低用户体验？", "会牺牲一部分覆盖率，但换来更高可信度。科研场景中可信度通常优先于流畅性。"),
        ],
    )

    doc.add_heading("四、关键配置项", level=1)
    add_table(
        doc,
        ["配置项", "默认值", "作用"],
        [
            ["RETRIEVAL_MODE", "hybrid", "vector、keyword、hybrid 三种检索模式"],
            ["HYBRID_CANDIDATE_K", "24", "混合检索阶段每路候选数量"],
            ["RRF_K", "60", "RRF 排名融合平滑参数"],
            ["ENABLE_RERANKER", "true", "是否启用轻量精排"],
            ["ENABLE_PARENT_CONTEXT", "true", "是否展开父块上下文"],
            ["STRICT_CITATION", "true", "是否对无引用回答触发拒答"],
            ["CHUNK_SIZE / CHUNK_OVERLAP", "800 / 120", "子 chunk 滑动窗口参数"],
        ],
        [2.0, 1.3, 3.2],
    )

    doc.add_heading("五、面试速记", level=1)
    add_bullets(
        doc,
        [
            "一句话讲项目：这是一个面向科研论文的证据约束型 RAG Agent，核心是结构化切分、混合召回、重排、引用约束和拒答。",
            "一句话讲父子分块：小块负责准确召回，大块负责完整回答。",
            "一句话讲混合检索：向量找语义，关键词找术语，RRF 负责融合排名。",
            "一句话讲 Reranker：召回要多而全，精排要少而准。",
            "一句话讲幻觉抑制：没有来源不回答，有来源也必须引用来源。",
        ],
    )

    doc.add_heading("六、当前边界与后续升级", level=1)
    add_bullets(
        doc,
        [
            "当前关键词检索是轻量本地实现，后续可替换为 BM25/Elasticsearch。",
            "当前 reranker 是规则打分，后续可接 bge-reranker、Cohere Rerank 或 DashScope rerank。",
            "当前查询改写是规则模板，后续可用 LLM 结合对话历史做指代消解和深度改写。",
            "当前评估以编译和轻量功能验证为主，后续可加入 HitRate、MRR、Faithfulness 数据集评估。",
            "多模态 RAG 仍建议作为独立增强链路，避免 MinerU 解析成本影响普通 RAG 稳定性。",
        ],
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_doc()
